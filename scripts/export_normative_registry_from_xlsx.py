"""Экспорт tmp/true_list_of_docs.xlsx → src/mr_norm/data/normative_documents_registry.json."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from mr_norm.data.normative_registry_key import (  # noqa: E402
    build_doc_key,
    format_adoption_date_display,
    format_adoption_date_iso,
    parse_date_cell,
)

DEFAULT_XLSX = ROOT / "tmp" / "true_list_of_docs.xlsx"
DEFAULT_OUT = ROOT / "src" / "mr_norm" / "data" / "normative_documents_registry.json"
SCHEMA_VERSION = 1


def _is_pue_row(doc_type: str, title: str, filename: str) -> bool:
    blob = f"{doc_type} {title} {filename}".casefold()
    return "пуэ" in blob or "правила устройства электроустановок" in blob


def export_registry(xlsx_path: Path, out_path: Path) -> int:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    documents: list[dict] = []
    seen_keys: set[str] = set()

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 6:
            continue
        filename_cell, doc_type, authority, doc_number, adoption_raw, title = row[:6]
        key_cell = row[6] if len(row) > 6 else None
        short_cell = row[7] if len(row) > 7 else None

        doc_type_s = str(doc_type or "").strip()
        title_s = str(title or "").strip()
        filename_s = str(filename_cell or "").strip()
        if not doc_type_s and not title_s:
            continue

        parsed = parse_date_cell(adoption_raw)
        registry_key = (
            str(key_cell).strip()
            if key_cell is not None and str(key_cell).strip()
            else build_doc_key(doc_type_s, doc_number, adoption_raw)
        )
        if registry_key in seen_keys:
            registry_key = f"{registry_key}_dup{len(seen_keys)}"
        seen_keys.add(registry_key)

        is_pue = _is_pue_row(doc_type_s, title_s, filename_s)
        match_stems: list[str] = []
        if filename_s and not is_pue:
            match_stems.append(Path(filename_s).stem)

        entry: dict = {
            "registry_key": registry_key,
            "doc_type": doc_type_s,
            "authority": str(authority or "").strip(),
            "doc_number": str(doc_number if doc_number is not None else "").strip(),
            "adoption_date": format_adoption_date_iso(parsed),
            "adoption_date_display": format_adoption_date_display(
                parsed, str(adoption_raw or "").strip()
            ),
            "title": title_s,
            "short_title": str(short_cell or "").strip(),
            "match_stems": match_stems,
        }
        if is_pue:
            entry["is_pue_canonical"] = True
        documents.append(entry)

    documents.sort(key=lambda d: (d.get("registry_key") or ""))
    payload = {"schema_version": SCHEMA_VERSION, "documents": documents}
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return len(documents)


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(description="Export normative registry from Excel")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()
    n = export_registry(args.xlsx, args.out)
    print(f"Wrote {args.out} ({n} documents)")


if __name__ == "__main__":
    main()
